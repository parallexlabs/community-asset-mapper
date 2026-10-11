"""Safe Excel parsing (read-only, data_only, no macros)."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from community_asset_mapper.security.io import MAX_FILE_BYTES


def read_excel(
    path: Path,
    sheet: str | None = None,
    max_bytes: int = MAX_FILE_BYTES,
) -> list[dict[str, str]]:
    """Read an Excel workbook sheet into string dictionaries."""
    resolved = path.resolve()
    size = resolved.stat().st_size
    if size > max_bytes:
        raise ValueError(f"Excel file exceeds size limit ({size} > {max_bytes})")
    workbook = load_workbook(resolved, read_only=True, data_only=True, keep_links=False)
    try:
        worksheet = workbook[sheet] if sheet else workbook.active
        if worksheet is None:
            return []
        rows_iter = worksheet.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if header_row is None:
            return []
        headers = [str(cell).strip() if cell is not None else "" for cell in header_row]
        rows: list[dict[str, str]] = []
        for row in rows_iter:
            record: dict[str, str] = {}
            for idx, header in enumerate(headers):
                if not header:
                    continue
                value = row[idx] if idx < len(row) else None
                record[header] = str(value).strip() if value is not None else ""
            rows.append(record)
        return rows
    finally:
        workbook.close()
